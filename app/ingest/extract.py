"""Document parsing.

Hand-written rather than using off-the-shelf loaders, because the stock loaders
drop exactly what matters most in real notes: embedded screenshots, speaker
notes, grouped shapes and the original paragraph/table ordering.
"""
import io
from pathlib import Path

import docx
import docx.table
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

from app.config import IMAGE_SUFFIXES, log, why
from app.ingest.vision import ImageCollector


class ExtractionError(ValueError):
    """Raised with a user-facing reason when a file cannot be parsed."""


def docx_text(data: bytes, images: ImageCollector) -> str:
    document = docx.Document(io.BytesIO(data))
    parts = []
    # Walk the body in order so paragraphs, tables and images keep their original sequence.
    for block in document.element.body.iterchildren():
        if block.tag == qn("w:p"):
            text = "".join(node.text or "" for node in block.iter(qn("w:t"))).strip()
            if text:
                parts.append(text)
            for blip in block.iter(qn("a:blip")):
                part = document.part.related_parts.get(blip.get(qn("r:embed")))
                token = images.add(part.blob) if part is not None else ""
                if token:
                    parts.append(token)
        elif block.tag == qn("w:tbl"):
            for row in docx.table.Table(block, document).rows:
                cells = []
                for cell in row.cells:
                    value = cell.text.strip()
                    # Merged cells repeat their text; keep one copy.
                    if value and (not cells or cells[-1] != value):
                        cells.append(value)
                if cells:
                    parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def shape_texts(shape, images: ImageCollector):
    """Yield text from a slide shape, recursing into groups."""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            yield from shape_texts(child, images)
        return
    try:
        blob = shape.image.blob
    except Exception:
        blob = None
    if blob:
        token = images.add(blob)
        if token:
            yield token
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            yield text
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                yield " | ".join(cells)


def pptx_pages(data: bytes, images: ImageCollector):
    presentation = Presentation(io.BytesIO(data))
    pages = []
    for number, slide in enumerate(presentation.slides, 1):
        parts = [text for shape in slide.shapes for text in shape_texts(shape, images)]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("备注：" + notes)
        pages.append((number, "\n".join(parts)))
    return pages


def pdf_pages(data: bytes, images: ImageCollector):
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for number, page in enumerate(reader.pages, 1):
        parts = [page.extract_text() or ""]
        try:
            page_images = list(page.images)
        except Exception as exc:
            # One unreadable image must not lose the whole page's text.
            log.info("pdf page %s images unreadable: %s", number, exc.__class__.__name__)
            page_images = []
        for pdf_image in page_images:
            try:
                token = images.add(pdf_image.image)
            except Exception:
                images.skipped += 1
                token = ""
            if token:
                parts.append(token)
        pages.append((number, "\n".join(parts)))
    return pages


def extract_pages(name: str, data: bytes, images: ImageCollector):
    """Return a list of (page_number_or_None, text); images become placeholder tokens."""
    suffix = Path(name).suffix.lower()

    if suffix in IMAGE_SUFFIXES:
        token = images.add(data)
        if not token:
            raise ExtractionError("图片无法读取，或尺寸太小（至少 100×100 像素）")
        return [(None, token)]

    if suffix in (".docx", ".pptx"):
        try:
            return ([(None, docx_text(data, images))] if suffix == ".docx"
                    else pptx_pages(data, images))
        except Exception as exc:
            log.warning("%s extraction failed for %s", suffix, name, exc_info=True)
            raise ExtractionError(
                f"{suffix} 解析失败：{why(exc)}（旧版 .doc/.ppt 请先另存为新格式）") from exc

    if suffix == ".pdf":
        try:
            return pdf_pages(data, images)
        except Exception as exc:
            log.warning("pdf extraction failed for %s", name, exc_info=True)
            raise ExtractionError(
                f"PDF 解析失败：{why(exc)}（加密或损坏的 PDF 需先解除保护）") from exc

    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return [(None, data.decode(encoding))]
        except UnicodeDecodeError:
            continue
    raise ExtractionError("无法识别文本编码，请保存为 UTF-8")
